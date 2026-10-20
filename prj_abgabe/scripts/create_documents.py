from pathlib import Path
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt as DocPt, RGBColor as DocColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
INK = "20272B"
RED = "C92C38"
TEAL = "007F86"
GRAY = "56646B"
LIGHT = "F2F5F6"
PINS = [
    ("CLK100MHZ", "E3", "Board-Oszillator", "100 MHz / 10 ns"),
    ("ck_rst", "C2", "RESET-Taster", "Aktiv Low"),
    ("btn_up", "D9", "BTN0", "Aktiv High; nur Phase 1"),
    ("seg_r[0] / a", "D13", "JA7 -> SSD AA", "Aktiv High"),
    ("seg_r[1] / b", "B18", "JA8 -> SSD AB", "Aktiv High"),
    ("seg_r[2] / c", "A18", "JA9 -> SSD AC", "Aktiv High"),
    ("seg_r[3] / d", "K16", "JA10 -> SSD AD", "Aktiv High"),
    ("seg_r[4] / e", "J17", "JB7 -> SSD AE", "Aktiv High"),
    ("seg_r[5] / f", "J18", "JB8 -> SSD AF", "Aktiv High"),
    ("seg_r[6] / g", "K15", "JB9 -> SSD AG", "Aktiv High"),
    ("dsel_r", "J15", "JB10 -> SSD C", "0: rechts; 1: links"),
]
PATTERNS = ["0111111", "0000110", "1011011", "1001111", "1100110",
            "1101101", "1111101", "0000111", "1111111", "1101111"]
SOURCES = [
    "Aufgabenstellung: prj_final/task/image3.png und image4.png (Phase I), image5.png (Phase II).",
    "Digilent Arty A7 Reference Manual: https://digilent.com/reference/programmable-logic/arty-a7/reference-manual",
    "Digilent Pmod SSD Reference Manual: https://digilent.com/reference/pmod/pmodssd/reference-manual",
    "Digilent Master-XDC: https://github.com/Digilent/digilent-xdc/blob/master/Arty-A7-35-Master.xdc",
    "GHDL-Dokumentation: https://ghdl.github.io/ghdl/",
    "Lokale Referenz: prj_final/src und prj_final/constraints/arty_a7_35_wecker.xdc. Pinbelegung und Anzeigemodule wurden daraus uebernommen.",
]


def section(title, paragraphs, table=None, code=None, image=None):
    return dict(title=title, paragraphs=paragraphs, table=table, code=code, image=image)


def font(size, bold=False):
    name = "calibrib.ttf" if bold else "calibri.ttf"
    path = Path("C:/Windows/Fonts") / name
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default(size=size)


def diagram(path, phase):
    image = Image.new("RGB", (1600, 650), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 25), f"PHASE {phase} | SIGNALFLUSS", font=font(36, True), fill=f"#{INK}")
    labels = (["BTN0", "2-FF-Sync +\nEntprellung", "Zaehler\n0 ... 9", "SSD-Treiber\nMux + Decoder"] if phase == 1 else
              ["100 MHz", "Sekunden-\nEnable", "Zaehler\n0 ... 99", "BCD + SSD\nMux + Decoder"])
    for index, label in enumerate(labels):
        left = 55 + index * 390
        draw.rounded_rectangle((left, 170, left + 310, 340), radius=8, fill=f"#{LIGHT}", outline=f"#{TEAL}", width=3)
        draw.multiline_text((left + 155, 255), label, font=font(30, True), fill=f"#{INK}", anchor="mm", align="center", spacing=8)
        if index < 3:
            draw.line((left + 310, 255, left + 380, 255), fill=f"#{INK}", width=4)
            draw.polygon([(left + 380, 255), (left + 365, 246), (left + 365, 264)], fill=f"#{INK}")
    draw.line((90, 465, 1480, 465), fill=f"#{RED}", width=4)
    draw.text((60, 520), "Gemeinsamer Systemtakt: 100 MHz | ck_rst: asynchron anfordern, synchron freigeben", font=font(28), fill=f"#{GRAY}")
    for index in range(1, 4):
        left = 55 + index * 390 + 155
        draw.line((left, 465, left, 355), fill=f"#{RED}", width=3)
    draw.text((60, 400), "Reset", font=font(26, True), fill=f"#{RED}")
    image.save(path)


def segment_image(path):
    image = Image.new("RGB", (1600, 650), "white")
    draw = ImageDraw.Draw(image)
    points = [(220, 110, 420, 110), (450, 140, 450, 290), (450, 340, 450, 490),
              (220, 520, 420, 520), (190, 340, 190, 490), (190, 140, 190, 290), (220, 315, 420, 315)]
    label_positions = [(320, 55), (510, 215), (510, 415), (320, 575), (125, 415), (125, 215), (320, 265)]
    for index, coordinates in enumerate(points):
        draw.line(coordinates, fill=f"#{RED}", width=24)
        draw.text(label_positions[index], "abcdefg"[index], font=font(36, True), fill=f"#{INK}", anchor="mm")
    draw.text((690, 130), "Bitfolge: g f e d c b a", font=font(44, True), fill=f"#{INK}")
    draw.text((690, 240), "0 = 0111111", font=font(44), fill=f"#{TEAL}")
    draw.text((690, 330), "1 = 0000110", font=font(44), fill=f"#{TEAL}")
    draw.text((690, 450), "1: Segment leuchtet", font=font(34), fill=f"#{GRAY}")
    draw.text((690, 510), "Gemeinsame Kathode / aktiv High", font=font(30), fill=f"#{GRAY}")
    image.save(path)


PHASE1 = [
    section("Ziel, Aufgabenstellung und Abgrenzung", [
        "Die erste Laborphase verbindet eine VHDL-Beschreibung mit einer sichtbaren Ausgabe am Arty A7-35T. Die Aufgabenbilder fordern die Ansteuerung eines Pmod SSD, eine Testbench und die Ausgabe der Zahl 0. Diese Basis wird hier um einen manuellen Dezimalzaehler erweitert: Jeder gueltige Druck auf BTN0 erhoeht die rechte Ziffer einmal. Nach 9 folgt 0; die linke Ziffer bleibt dunkel.",
        "Die Erweiterung veraendert das Lernziel nicht: Segmentcodierung, Pinzuordnung und zeitliches Verhalten bleiben direkt beobachtbar. Nach Einschalten beziehungsweise Reset steht die rechte Anzeige auf 0. Die zweite Phase nutzt denselben Anzeigepfad fuer einen Sekundenzaehler. Die dritte Phase ist eine getrennte, unveraenderte Kopie des vorhandenen Weckers.",
        "Dieser Bericht dokumentiert den implementierten Stand. Die automatisierte RTL-Simulation wurde mit GHDL 6.0.0 ausgefuehrt. Eine erfolgreiche Synthese, Timing Closure oder eine Messung am physischen Board wird nicht behauptet. Die entsprechenden Abnahmeschritte sind als offene Pruefliste aufgenommen.",
    ], table=(['ID', 'Anforderung', 'Nachweis'], [
        ('P1-01', 'Startwert 0, links dunkel', 'Testbench + Boardpruefung offen'),
        ('P1-02', 'Ein Inkrement pro Tastendruck', 'Prell-/Halte-/Loslasstest'),
        ('P1-03', 'Dezimaler Umlauf 0 bis 9', 'Alle zehn Muster simuliert'),
        ('P1-04', 'Reset ueber ck_rst', 'Reset nach Zaehlschritten simuliert'),
        ('P1-05', 'Eigenstaendiges Vivado-Projekt', 'Tcl + XDC enthalten; Build offen'),
    ])),
    section("Hardware und sichere Verdrahtung", [
        "Verwendet wird der Arty A7-35T mit dem FPGA-Part xc7a35ticsg324-1L und dem 100-MHz-Oszillator. Das externe Digilent Pmod SSD besitzt zwei Ziffern mit gemeinsam genutzten Segmentleitungen. Es handelt sich nicht um eine bereits auf dem Arty integrierte Anzeige. Ein Pmod SSD reicht fuer Phase I und II aus.",
        "Die Verdrahtung entspricht dem bestehenden Weckerprojekt: JA und JB werden jeweils in der UNTEREN Reihe verwendet. Die Segmentanschluesse AA bis AG und der Ziffernselektor C des SSD werden einzeln zugeordnet. Versorgung und gemeinsame Masse sind zusaetzlich erforderlich; die Signaltabelle ersetzt diese Versorgungsanschluesse nicht. Vor dem Umstecken die Versorgung ausschalten, Steckerausrichtung kontrollieren und nur die fuer das Modul vorgesehene 3,3-V-Versorgung verwenden.",
        "Alle Signale verwenden LVCMOS33. BTN0 ist aktiv High, der Reset-Eingang ck_rst aktiv Low. Die Annahme dsel_r=0 fuer die rechte Ziffer stammt aus dem vorhandenen ssd_driver. Die reale Zuordnung wird beim Boardtest mit einer bekannten Ziffer kontrolliert.",
    ], table=(['Port / Segment', 'FPGA-Pin', 'Verbindung', 'Bedeutung'], PINS)),
    section("Theorie der 7-Segment-Anzeige", [
        "Sieben einzeln ansteuerbare Leuchtsegmente a bis g bilden eine Dezimalziffer. Die gemeinsame Kathode und die Beschaltung des Pmod SSD fuehren zu aktiven High-Signalen: Eine logische 1 schaltet das betreffende Segment ein. Im VHDL-Vektor seg(6 downto 0) steht Bit 6 fuer g und Bit 0 fuer a. Diese Reihenfolge ist bei jeder Fehlersuche zuerst zu kontrollieren.",
        "Der Decoder ist kombinatorisch. Er uebersetzt eine 4-Bit-Ziffer unmittelbar in ein Segmentmuster und benoetigt keinen eigenen Takt. Fuer Werte ausserhalb 0 bis 9 zeigt der uebernommene Decoder einen Mittelstrich; blank=1 schaltet alle Segmente ab. Im normalen Zaehlerbetrieb werden ausschliesslich gueltige Dezimalwerte erzeugt.",
    ], table=(['Ziffer', 'gfedcba', 'Leuchtende Segmente'], [
        (str(value), pattern, ', '.join(letter for bit, letter in enumerate('abcdefg') if pattern[6-bit] == '1'))
        for value, pattern in enumerate(PATTERNS)
    ]), image='segments.png'),
    section("Entwurf und Gesamtsystem Phase I + II", [
        "In Phase I durchlaeuft BTN0 zuerst zwei Synchronisationsregister und danach eine Stabilitaetspruefung. Nur der akzeptierte Wechsel von Low nach High veraendert den Zaehler. Der Zaehler stellt eine Ziffer bereit; ssd_driver waehlt die aktive Stelle und seg7_decoder erzeugt die Segmentbits. Der linke Multiplex-Slot wird durch Nullsetzen der Ausgaenge ausgeblendet.",
        "Phase II ersetzt den manuellen Ereignispfad durch einen synchronen Sekundenteiler. Ein Zaehler von 0 bis 99 liefert Einer und Zehner an denselben Treiber. Der Systemtakt bleibt in beiden Phasen unveraendert; Zaehlen erfolgt ueber eine Bedingung innerhalb des getakteten Prozesses und nicht ueber einen neu erzeugten langsamen Takt.",
        "Die gemeinsame Resetstruktur setzt die Steuerung und den Anzeigezustand zurueck. Das aktive Low-Eingangssignal fordert Reset asynchron an; zwei Register geben ihn synchron zum Systemtakt frei. Die eigentlichen Zustandsregister reagieren am naechsten Systemtakt. Alle neuen Module sind bewusst klein gehalten, damit der Datenfluss im Labor nachvollzogen werden kann.",
    ], image='block_phase1.png', table=(['Baustein', 'Aufgabe'], [
        ('phase1_top / phase2_top', 'Eingang, Reset und Zaehlersteuerung'),
        ('ssd_driver', 'Stellenauswahl und zeitliches Multiplexen'),
        ('seg7_decoder', 'Dezimalziffer zu gfedcba'),
        ('tick_gen / clock_pkg', 'Refresh-Enable und Hilfsfunktion max_nat'),
    ])),
    section("Synchronisation, Entprellung und Reset", [
        "Ein mechanischer Taster ist zum FPGA-Takt asynchron. Zwei hintereinander geschaltete Flipflops reduzieren die Wahrscheinlichkeit, dass ein metastabiler Zustand in die nachfolgende Logik gelangt. Sie ersetzen keine Entprellung. Eine reine Flankenerkennung am ungefilterten Signal koennte waehrend eines einzigen Tastendrucks mehrere Inkremente ausloesen.",
        "Der Entprellzaehler prueft daher, ob der synchronisierte Eingang vom bisher akzeptierten Zustand abweicht. Sobald der Eingang zum alten Zustand zurueckkehrt, beginnt die Wartezeit erneut. Erst DEBOUNCE_CYCLES aufeinanderfolgende abweichende Abtastungen uebernehmen den neuen Zustand. Bei 100 MHz und 1.000.000 Takten betraegt das Fenster 10 ms, zuzueglich der Synchronisationslatenz von wenigen Systemtakten.",
        "Nur bei der Uebernahme eines High-Zustands wird inkrementiert. Ein dauerhaft gehaltener Taster erzeugt keine weiteren Ereignisse. Auch das Loslassen muss stabil sein, bevor ein neuer Druck erkannt werden kann. Anders als die Entprellung des Weckers besitzt diese Implementierung absichtlich keine Auto-Repeat-Funktion.",
        "ck_rst setzt die Reset-Synchronisationskette asynchron auf 11. Nach Freigabe wird sie ueber zwei steigende Flanken mit Nullen gefuellt. Die Fachlogik sieht den freigegebenen Zustand danach. Ein beim Freigeben gehaltener BTN0 wird nach der Entprellzeit als neuer Druck akzeptiert. Reset dominiert dabei stets die Zaehleroperation.",
    ], table=(['Parameter / Zustand', 'Wert / Wirkung'], [
        ('Systemtakt', '100.000.000 Hz'), ('Entprellfenster', '1.000.000 Takte = 10 ms'),
        ('stable=0 -> stable=1', 'Genau ein Inkrement'), ('stable bleibt 1', 'Keine Wiederholung'),
        ('stable=1 -> stable=0', 'Nur Freigabe fuer naechsten Druck'),
    ])),
    section("VHDL-Implementierung des manuellen Zaehlers", [
        "Die Datei src/phase1_top.vhd enthaelt Resetkette, Eingangssynchronisation, Entprellzustand und Zaehler. Der Zahlenbereich natural range 0 to 9 beschreibt explizit den gueltigen Zustand. Vor dem Inkrement wird der Grenzwert geprueft, damit der Zaehler nie den zulaessigen Bereich verlaesst.",
        "Der folgende Ausschnitt zeigt die Entscheidung nach abgeschlossenem Entprellfenster. Die umgebende Logik setzt debounce_count bei gleichen Abtastwerten zurueck und erhoeht ihn andernfalls. Es gibt weder wait-Anweisungen noch Simulationsverzoegerungen im synthetisierbaren Teil. Die vollstaendigen Quellen liegen im Ordner src; die zugehoerige Testbench liegt unter sim.",
        "count wird mit to_unsigned(count, 4) an den Anzeigetreiber uebergeben. Die zweite Ziffer erhaelt den Wert 0, ihr aktiver Slot wird jedoch am Ausgang ausgeblendet. Dadurch entsteht keine sichtbare fuehrende Null. Alle vier uebernommenen Basismodule sind lokale Kopien, sodass das Projekt ohne Verweise auf prj_final geoeffnet werden kann.",
    ], code="""elsif debounce_count = DEBOUNCE_CYCLES - 1 then
    debounce_count <= 0;
    stable <= button_sync(1);
    if button_sync(1) = '1' then
        if count = 9 then
            count <= 0;
        else
            count <= count + 1;
        end if;
    end if;
else
    debounce_count <= debounce_count + 1;
end if;""", table=(['Datei', 'Rolle'], [
        ('src/phase1_top.vhd', 'Top-Level des Tasterzaehlers'),
        ('src/ssd_driver.vhd', 'Unveraenderte Uebernahme aus prj_final'),
        ('src/seg7_decoder.vhd', 'Unveraenderte Uebernahme aus prj_final'),
        ('src/tick_gen.vhd, clock_pkg.vhd', 'Unveraenderte Basismodule'),
        ('sim/tb_phase1_top.vhd', 'Selbstpruefende Testbench mit Assertions'),
    ])),
    section("Zeitmultiplex und wahrgenommene Anzeige", [
        "Ein Pmod SSD nutzt sieben gemeinsame Segmentleitungen und ein Auswahlsignal fuer zwei Ziffern. Die FPGA-Logik schaltet deshalb regelmaessig zwischen den Stellen um. Zu jedem Slot muessen Segmentmuster und Stellenauswahl zusammenpassen. Phase I aktiviert nur den rechten Slot; Phase II zeigt in beiden Slots unterschiedliche Dezimalstellen.",
        "Der uebernommene Treiber berechnet DIV = max(1, CLK_HZ / (2 * REFRESH_HZ)). Mit 100 MHz und REFRESH_HZ=500 ergibt sich DIV=100.000. Die Stellenauswahl wechselt alle 1 ms; ein vollstaendiger Durchlauf beider Stellen dauert 2 ms. Jede Ziffer wird damit 500-mal pro Sekunde mit ungefaehr 50 Prozent Tastverhaeltnis angesteuert.",
        "Die schnelle Wiederholung laesst die Anzeige fuer das Auge ruhig erscheinen. Die Entprellung von 10 ms und der spaetere Sekundenzaehler sind davon unabhaengig. Helligkeit, elektrische Umschalteffekte und moegliches Ghosting sind Eigenschaften, die an der realen Anzeige kontrolliert werden muessen; eine digitale RTL-Simulation bewertet keine analogen Leuchtstaerken.",
        "Bei der Fehlersuche zuerst bekannte Ziffern 0, 1 und 8 einstellen. Eine falsche Segmentreihenfolge fuehrt zu reproduzierbar falschen Mustern; vertauschte Ziffern deuten dagegen auf die Select-Zuordnung. Unruhiges Zaehlen bei unveraendertem Anzeigebild weist eher auf den Eingangspfad als auf den Multiplexer hin.",
    ], table=(['Groesse', 'Wert'], [
        ('Systemperiode', '10 ns'), ('Halbbild / Slot', '1 ms'),
        ('Vollstaendiges Bild', '2 ms'), ('Wiederholrate je Ziffer', '500 Hz'),
        ('Dunkler Slot in Phase I', 'dsel_r=1; seg_r=0000000'),
    ])),
    section("Pin-Constraints und Vivado-Ablauf", [
        "constraints/arty_a7_35_ssd.xdc definiert die physikalischen Pins und LVCMOS33. create_clock beschreibt die reale Eingangsperiode von 10 ns. Der optionale Port btn_up wird nur dann zugeordnet, wenn er im Top-Level vorhanden ist; dieselbe Datei kann dadurch auch fuer Phase II verwendet werden. Die Zuordnung ist aus dem Wecker-XDC abgeleitet und auf die benoetigten Signale reduziert.",
        "Das Skript scripts/create_project.tcl waehlt den Top-Level anhand des Ordnernamens phase_1 beziehungsweise phase_2. Es legt ein lokales Projekt unter vivado an, fuegt Quellen, Testbench und Constraints hinzu und aktiviert VHDL-2008. Dieser Standard ist insbesondere fuer std.env.finish in den Testbenches erforderlich. Das Projekt muss aus einem dieser beiden Ordner erzeugt werden.",
        "Im Vivado-Tcl-Terminal kann source scripts/create_project.tcl ausgefuehrt werden. Alternativ dient der folgende Aufruf aus dem Phasenordner. Danach Behavioral Simulation starten, anschliessend Synthesis, Implementation und Generate Bitstream ausfuehren. Vor dem Programmieren sind DRC, nicht zugeordnete Ports und Setup-/Hold-Timing zu kontrollieren. Es liegen noch keine Vivado-Ergebnisberichte vor.",
    ], code="""vivado -mode gui -source scripts/create_project.tcl

# GHDL-Alternative aus dem Phasenordner:
./scripts/run_sim.ps1 -Ghdl ghdl
# Bei nicht gesetztem PATH einen absoluten ghdl.exe-Pfad angeben.""", table=(['Pruefpunkt', 'Akzeptanzkriterium'], [
        ('Top-Level', 'phase1_top / tb_phase1_top'),
        ('Part', 'xc7a35ticsg324-1L'), ('IO-Standard', 'Alle externen Signale LVCMOS33'),
        ('Timing', '10-ns-Systemtakt; keine negativen Slacks'),
        ('Hardware Manager', 'Passendes Board und eigener Bitstream'),
    ])),
    section("Simulation und nachgewiesene Ergebnisse", [
        "tb_phase1_top ist selbstpruefend: Assertions vergleichen die sichtbaren Segmentausgaenge mit unabhaengig definierten Dezimalmustern. Jeder Vergleich erfolgt an der fallenden Taktflanke, nachdem die getaktete Logik und die kombinatorische Decodierung eingeschwungen sind. Je Ziffer werden 24 Takte beobachtet; damit werden bei den Testparametern beide Multiplex-Slots erreicht.",
        "Die Simulation setzt CLK_HZ=100, DEBOUNCE_CYCLES=4 und REFRESH_HZ=10. Das sind bewusst beschleunigte Funktionsparameter; der Testtakt selbst hat 10 ns Periode. Es wird damit nicht behauptet, dass die Simulation echte 10 ms oder das elektrische Prellen eines konkreten Tasters nachbildet. Geprueft wird die Logik fuer stabile und zu kurze Eingabesequenzen.",
        "Ausgefuehrt mit GHDL 6.0.0: PASS: phase 1 startup, bounce, hold, release, wrap, reset. Die Testbench beendet sich nach 7.290 ns ohne Assertion-Fehler. Das Skript run_sim.ps1 erzeugt eine simulation.ghw unter sim/build; ein abweichender GHDL-Pfad kann als Parameter uebergeben werden.",
    ], table=(['Test', 'Stimulus', 'Beobachtetes Ergebnis'], [
        ('Start / Reset', 'Reset aktiv, danach Freigabe', 'Rechte 0, links aus'),
        ('Prellen', 'Mehrere 1-Takt-Wechsel', 'Kein Inkrement'),
        ('Druck', 'High fuer mindestens 4 Abtastungen', '0 -> 1'),
        ('Halten', 'Weitere 100 Takte High', 'Bleibt 1'),
        ('Loslassen', 'Stabiles Low', 'Kein Inkrement'),
        ('Alle Ziffern', 'Weitere gueltige Druecke', 'Muster 0 bis 9 korrekt'),
        ('Ueberlauf', 'Gueltiger Druck bei 9', '9 -> 0'),
        ('Erneuter Reset', 'Reset nach weiterem Druck', 'Zurueck auf 0'),
    ])),
    section("Boardabnahme, Fazit und Quellen", [
        "Die logischen Anforderungen sind in der RTL-Simulation erfuellt. Offen bleiben Synthese, Implementation, Timing-Analyse und der physische Boardtest. Die folgende Liste ist deshalb ein Abnahmeplan und kein ausgefuelltes Messprotokoll. Vor der endgueltigen Abgabe sollen Datum, verwendetes Board, Vivado-Version und tatsaechliche Beobachtungen ergaenzt werden.",
        "Board-Pruefliste: [ ] Verdrahtung bei ausgeschalteter Versorgung kontrolliert; [ ] Reset zeigt rechts 0 und links dunkel; [ ] zehn einzelne Tastendruecke ergeben einen Umlauf; [ ] langes Halten zaehlt nur einmal; [ ] Loslassen zaehlt nicht; [ ] Reset funktioniert aus beliebiger Ziffer; [ ] keine sichtbaren falschen Segmente oder stoerendes Flackern; [ ] Timing-Bericht ohne Verletzungen abgelegt.",
        "Fazit: Die Phase trennt asynchrone Eingabebehandlung, Zustandsaenderung und Anzeige. Das vereinfacht die Fehlersuche und ermoeglicht die Weiterverwendung des Anzeigepfads in Phase II. Der entscheidende Unterschied zur spaeteren Stoppuhr ist ausschliesslich die Quelle des Inkrementereignisses.",
        "Anhang / Artefakte: Vollstaendige VHDL-Quellen unter src/, selbstpruefende Testbench unter sim/, Pin-Constraints unter constraints/ und reproduzierbare Aufrufe unter scripts/. Die PowerPoint-Datei enthaelt Sprechernotizen; die Praesentation ist auf etwa 20 Minuten ohne Fragerunde ausgelegt. Autor und Gruppe sind vor der Abgabe persoenlich zu ergaenzen.",
        *SOURCES,
    ]),
]


PHASE2 = [
    section("Phase II: Anforderungen und Bedienung", [
        "Die zweite Phase erweitert die Anzeige zu einem automatischen Dezimalzaehler. Gemaess Aufgabenbild image5.png wird von 0 bis 99 im Sekundentakt gezaehlt und ein Reset per Taster vorgesehen. Die Anzeige verwendet beide Stellen und stellt den Startwert als 00 dar. Nach 99 folgt 00; der Umlauf dauert nominell 100 Sekunden.",
        "Es gibt bewusst keinen Start-/Stop-Taster: Nach Freigabe von ck_rst laeuft der Zaehler automatisch. Ein gedrueckter Reset haelt sowohl den Zahlenstand als auch den Sekundenteiler auf null. Nach dem Loslassen beginnt ein neues vollstaendiges Sekundenintervall. BTN0 wird in Phase II nicht benoetigt und ist kein Top-Level-Port.",
        "Die ersten zehn Seiten dieses kombinierten Berichts dokumentieren die wiederverwendete Basis aus Phase I. Die folgenden zehn Seiten behandeln die Erweiterung. Der Hardwareaufbau fuer das Pmod SSD bleibt unveraendert. Phase III gehoert nicht zur Sekundenzaehler-Implementierung und liegt separat als Weckerkopie vor.",
    ], table=(['ID', 'Anforderung', 'Akzeptanzkriterium'], [
        ('P2-01', 'Anzeige 00 bis 99', 'Einer und Zehner korrekt'),
        ('P2-02', 'Ein Schritt pro Sekunde', '100.000.000 Systemtakte je Schritt'),
        ('P2-03', 'Umlauf', '99 -> 00 ohne Zwischenwert 100'),
        ('P2-04', 'Reset per Taster', 'Zahl und Teiler werden geloescht'),
        ('P2-05', 'Wiederverwendung', 'Unveraenderte Anzeigemodule aus Phase I'),
    ])),
    section("Sekundenbasis ohne zusaetzliche Taktdomaene", [
        "Der FPGA arbeitet weiterhin ausschliesslich mit 100 MHz. Ein synchroner Teiler zaehlt von 0 bis CLK_HZ-1. Am Endwert setzt er sich auf 0 und erlaubt im selben Prozess genau ein Inkrement des Dezimalzaehlers. Diese Bedingung wirkt wie ein Clock-Enable. Es wird keine langsame Taktleitung im Logiknetz erzeugt.",
        "N = f_clk / f_step = 100.000.000 / 1 = 100.000.000. Die Schrittzeit betraegt T_step = N * 10 ns = 1 s. Der Endwert ist 99.999.999, nicht 100.000.000; andernfalls entstuende ein Off-by-one-Fehler. Fuer die Darstellung des Teilers sind mindestens 27 Bits erforderlich, weil 2^26 < 100.000.000 <= 2^27 gilt.",
        "CLK_HZ ist ein positiver Generic und entspricht auf dem Board der realen Taktfrequenz. In der Simulation darf er kleiner gewaehlt werden, damit ein voller Durchlauf schnell geprueft werden kann. Die reale Zeitgenauigkeit bleibt an die Toleranz des Board-Oszillators gebunden; das Design enthaelt keine externe Referenz oder Kalibrierung.",
        "Der erste Schritt erfolgt nach dem ersten vollstaendigen Intervall nach interner Resetfreigabe. Die beiden Reset-Synchronisationsregister verursachen beim Loslassen eine kurze Freigabelatenz im Bereich weniger 10-ns-Takte. Diese ist gegenueber einer Sekunde vernachlaessigbar, wird in der Testbench aber explizit beruecksichtigt.",
    ], table=(['Groesse', 'Hardware', 'Beschleunigte Simulation'], [
        ('CLK_HZ', '100.000.000', '1.000'), ('Teiler-Endwert', '99.999.999', '999'),
        ('Takte je Schritt', '100.000.000', '1.000'),
        ('Physische Zeit je Schritt', '1 s', '10 us bei 10-ns-Testtakt'),
    ])),
    section("Dezimalzaehler und Stellenzerlegung", [
        "Der Zahlenstand wird als natural range 0 to 99 gespeichert. Das spart eine separate Uebertragssteuerung zwischen zwei BCD-Registern und formuliert den zulaessigen Zustand unmittelbar. Falls count=99 und der Sekundenteiler ausloest, wird count auf 0 gesetzt. Ansonsten wird count um eins erhoeht.",
        "Fuer die Anzeige wird count kombinatorisch zerlegt: ones = count mod 10 und tens = count / 10. Beide Ergebnisse liegen von 0 bis 9 und werden mit to_unsigned(..., 4) in die Eingangsvektoren des SSD-Treibers umgewandelt. Die Division durch eine konstante Zehn beschreibt synthetisierbare kombinatorische Logik, keine Software-CPU und keine neue Taktdomaene.",
        "Beim Wechsel 09 -> 10 aendern sich beide Stellen aus demselben Zaehlerregister. Der Multiplexer zeigt immer die gerade ausgewaehlte Stelle des aktuellen Zustands. Die Testbench vergleicht deshalb in jedem Slot nur die selektierte Ziffer, statt die beiden Stellen als gleichzeitig separat verdrahtete Ausgaenge zu behandeln.",
        "Die fuehrende Null bleibt absichtlich sichtbar: 00, 01, ..., 09, 10, ..., 99. So laesst sich der zweistellige Bereich eindeutig ablesen und der Uebergang zwischen Einer- und Zehnerstelle im Labor pruefen.",
    ], table=(['count', 'count / 10', 'count mod 10', 'Anzeige'], [
        ('0', '0', '0', '00'), ('9', '0', '9', '09'), ('10', '1', '0', '10'),
        ('59', '5', '9', '59'), ('99', '9', '9', '99'), ('nach 99', '0', '0', '00'),
    ])),
    section("Architektur und Wiederverwendung", [
        "phase2_top besitzt vier Ports: CLK100MHZ, ck_rst, seg_r und dsel_r. Der interne Datenpfad besteht aus Sekundenteiler, Zahlenregister, Stellenzerlegung und ssd_driver. Der Refresh-Pfad des Treibers arbeitet unabhaengig vom Sekundenereignis. Dadurch bleibt die Anzeige auch waehrend eines unveraenderten Zahlenstands aktiv.",
        "Die Dateien clock_pkg.vhd, tick_gen.vhd, seg7_decoder.vhd und ssd_driver.vhd sind identische lokale Kopien aus prj_final. Die Phase haengt somit nicht von relativen Dateiverweisen in andere Projekte ab. Das erleichtert die Abgabe und das Oeffnen auf einem anderen Rechner.",
        "Der manuelle Eingangspfad aus Phase I wird nicht instanziiert, weil Phase II keinen Zaehltaster benoetigt. Gemeinsam ist die Resetstruktur mit asynchroner Anforderung und synchroner Freigabe. Sie setzt sowohl die Zeitbasis als auch den Zahlenstand zurueck. Ein Reset waehrend eines fast vollstaendigen Intervalls darf keinen vorzeitigen Folgeschritt verursachen.",
    ], image='block_phase2.png', table=(['Aspekt', 'Phase I', 'Phase II'], [
        ('Ereignisquelle', 'Entprellter BTN0', 'Sekundenteiler'),
        ('Zahlenbereich', '0 bis 9', '0 bis 99'), ('Stellen', 'Rechts; links dunkel', 'Beide aktiv'),
        ('Reset', 'Zahl und Eingabezustand', 'Zahl und Sekundenphase'),
        ('Anzeige-Basis', 'ssd_driver + Decoder', 'Unveraendert wiederverwendet'),
    ])),
    section("RTL-Kern und Resetprioritaet", [
        "Der folgende Ausschnitt ist der Kern des getakteten Prozesses in phase2_top. Reset hat hoechste Prioritaet. Am Teiler-Endwert werden Teiler und Zahlenstand gemeinsam aktualisiert; an allen anderen Takten aendert sich nur der Teiler. Damit gibt es genau eine Quelle fuer jede Zustandszuweisung.",
        "Die Zuweisungen sind VHDL-Signalzuweisungen. Vergleiche lesen innerhalb eines Taktaufrufs den bisherigen Registerzustand; die neuen Werte werden danach wirksam. Die Endwertpruefung auf CLK_HZ-1 ist genau auf dieses Verhalten abgestimmt. Ein separater registrierter Tick im Zaehlerpfad waere ebenfalls moeglich, wuerde hier aber nur zusaetzliche Latenz und Erklaerungsaufwand erzeugen.",
        "Beim Reset werden divider und count gleichzeitig null. Das ist fuer reproduzierbares Verhalten wichtiger als nur eine geloeschte Anzeige: Ein stehen gebliebener Teilerrest koennte sonst unmittelbar nach dem Loslassen einen verkuerzten ersten Zaehlschritt ausloesen. Die Testbench prueft diese Eigenschaft mit einem Reset mitten im Intervall.",
    ], code="""if reset_sync(1) = '1' then
    divider <= 0;
    count <= 0;
elsif divider = CLK_HZ - 1 then
    divider <= 0;
    if count = 99 then
        count <= 0;
    else
        count <= count + 1;
    end if;
else
    divider <= divider + 1;
end if;

ones <= to_unsigned(count mod 10, 4);
tens <= to_unsigned(count / 10, 4);"""),
    section("Testmethodik und beschleunigte Zeitbasis", [
        "Ein kompletter Durchlauf bei realen Parametern wuerde zehn Milliarden Systemtakte benoetigen. Deshalb instanziiert tb_phase2_top das DUT mit CLK_HZ=1.000 und REFRESH_HZ=100. Der Testtakt bleibt 10 ns lang. Ein Zaehlschritt benoetigt damit 1.000 Takte beziehungsweise 10 us Simulationszeit, waehrend die Registerlogik und die Umlaufregeln unveraendert bleiben.",
        "Nach Resetfreigabe wartet die Testbench genau zwei Takte fuer die Freigabekette. Anschliessend berechnet sie den Sollstand aus der seitdem vergangenen Taktzahl: expected = (elapsed / 1.000) mod 100. Je nach dsel_r wird die erwartete Einer- oder Zehnerziffer in ein unabhaengig hinterlegtes Segmentmuster uebersetzt. Verglichen wird an jeder fallenden Flanke.",
        "Die erste Laufstrecke umfasst 101.357 Takte und damit mehr als einen vollstaendigen Umlauf. Danach erfolgt Reset bei nicht vollstaendig abgelaufenem Sekundenintervall. Eine zweite Strecke ueber 2.000 Takte prueft den Neustart. Zusaetzliche Flags bestaetigen, dass beide Multiplex-Slots tatsaechlich beobachtet wurden.",
        "Diese Simulation beweist das digitale Verhalten fuer die ausgefuehrten Stimuli, nicht die analoge Entprellqualitaet, die elektrische LED-Ansteuerung oder das Hardware-Timing nach Place-and-Route. Solche Aussagen erfordern den im spaeteren Abschnitt beschriebenen Board- und Timingnachweis.",
    ], code="""./scripts/run_sim.ps1 -Ghdl ghdl

# Artefakte:
# sim/build/work-obj08.cf   GHDL-Arbeitsbibliothek
# sim/build/simulation.ghw  Wellenform
# sim/testergebnis.txt      protokollierter Testlauf"""),
    section("Testergebnisse und Anforderungsabdeckung", [
        "Der Testlauf mit GHDL 6.0.0 beendet sich nach 1.033.690 ns mit PASS: phase 2 exact period, all 00..99, 09->10, 99->00, mux, reset/restart. Es wurde kein Assertion-Fehler ausgeloest. Der Test kontrolliert kontinuierlich die segmentcodierte Ausgabe und nicht nur einen internen Zaehlerwert.",
        "Insbesondere werden zu fruehe oder zu spaete Zaehlschritte erkannt, weil der erwartete Stand an jeder Taktgrenze aus der verstrichenen Zeit abgeleitet wird. Der komplette Durchlauf beinhaltet alle Dezimalstellen und alle Zehneruebertraege. Nach dem Reset ist der Sollzeitpunkt des ersten neuen Inkrements wieder exakt an die neue Intervallgrenze gebunden.",
        "Auch Phase I wurde nach Erstellung der gemeinsamen Skripte erneut simuliert und bleibt erfolgreich. Die Weckerkopie wurde dagegen nicht neu entworfen; ihre 29 Dateien wurden bytegleich per SHA-256-Vergleich geprueft. Daraus folgt keine neue Funktionsverifikation des Weckers.",
    ], table=(['Test / Anforderung', 'Pruefung', 'Status'], [
        ('00 beim Start', 'Resetwert am Segmentausgang', 'Bestanden'),
        ('Exakte Schrittperiode', 'Jeder der 1.000 Testtakte je Schritt', 'Bestanden'),
        ('09 -> 10', 'BCD-Wechsel im Gesamtumlauf', 'Bestanden'),
        ('00 bis 99', 'Alle Zustaende mit Slotvergleich', 'Bestanden'),
        ('99 -> 00', 'Erster voller Umlauf', 'Bestanden'),
        ('Beide Ziffern', 'Beobachtungsflags im Test', 'Bestanden'),
        ('Reset / Intervallneustart', 'Reset nach 101.357 Takten', 'Bestanden'),
        ('Board / Synthese / Timing', 'Erfordert Vivado und Hardware', 'Offen'),
    ])),
    section("Synthese, Timing und Inbetriebnahme", [
        "Die Projektanlage erfolgt analog zu Phase I ueber scripts/create_project.tcl, nun mit phase2_top und tb_phase2_top. Der ausgewaehlte FPGA-Part muss zum realen Board passen. Bei einem anderen Arty-Modell sind Part und Pin-Constraints vor der Implementierung zu pruefen; die vorliegende Abgabe ist fuer das A7-35T vorbereitet.",
        "Der Systemtakt ist die einzige funktionale Taktdomaene. Trotzdem muessen im implementierten Design alle relevanten Datenpfade die 10-ns-Vorgabe einhalten. Besonders die konstante Division/Modulo-Logik und der breite Teiler sind im Synthesebericht zu betrachten. Es werden keine Ressourcen- oder Slackwerte erfunden: Ein konkreter Vivado-Lauf wurde hier mangels verfuegbarem Vivado nicht ausgefuehrt.",
        "Nach Behavioral Simulation folgen Synthesis, Implementation, DRC und Timing Summary. Erst bei fehlerfreier Pinzuordnung und erfolgreichen Timingpruefungen wird ein Bitstream erzeugt und ueber den Hardware Manager programmiert. Der Reset-Taster soll die Anzeige auf 00 halten. Nach Freigabe folgt 01 erst nach etwa einer Sekunde.",
        "Eine Stoppuhrmessung ueber mehrere Intervalle ist aussagekraeftiger als die subjektive Bewertung eines einzelnen Schritts. Fuer eine belastbare Frequenzmessung sind Messaufbau, Aufloesung und Beobachtungsdauer zu dokumentieren. Ohne diese Angaben ist die Aussage 'genau eine Sekunde' nur eine nominelle Rechnung auf Basis des 100-MHz-Eingangstakts.",
    ], table=(['Zu protokollieren', 'Aktueller Stand'], [
        ('Vivado-Version / Datum', 'Noch einzutragen'),
        ('LUT / FF / IO / Clock-Ressourcen', 'Noch nicht gemessen'),
        ('Worst Setup Slack / Hold Slack', 'Noch nicht gemessen'),
        ('DRC / unzugeordnete Pins', 'Noch nicht ausgefuehrt'),
        ('Bitstream / Boardkennung', 'Noch einzutragen'),
    ])),
    section("Gemeinsamer Boardtest und Fehlerdiagnose", [
        "Der gemeinsame Hardwaretest beginnt mit Phase I, weil sich damit jedes Segmentmuster manuell anhalten laesst. Erst wenn 0, 1 und 8 korrekt erscheinen und der Taster einmalig zaehlt, wird Phase II programmiert. So werden Verdrahtungsfehler von Fehlern in der Zeitbasis getrennt.",
        "Phase-II-Pruefliste: [ ] Startwert 00; [ ] erster Schritt nach einem vollen Intervall; [ ] 09 -> 10 korrekt; [ ] 99 -> 00 nach nominell 100 s; [ ] Reset bei laufendem Zaehler; [ ] Reset laenger halten; [ ] nach Freigabe wieder volles Intervall; [ ] beide Ziffern ohne stoerendes Flackern; [ ] alle Ergebnisse mit Datum und Aufbau dokumentiert.",
        "Die nachfolgende Diagnose ist als Hypothesenliste gedacht. Aenderungen duerfen erst nach einer gezielten Pruefung vorgenommen werden. Besonders bei scheinbar vertauschten Ziffern nicht gleichzeitig Segmente und Select-Signal umbelegen, weil dadurch eine eindeutige Fehlerursache verloren ginge.",
    ], table=(['Symptom', 'Naechste gezielte Pruefung'], [
        ('Alles dunkel', 'Versorgung, Masse, Bitstream und Resetpegel'),
        ('Falsche Segmente', 'gfedcba und AA..AG gegen Pinliste pruefen'),
        ('Ziffern vertauscht', 'dsel_r / SSD C mit bekannten Werten testen'),
        ('Zaehler zu schnell / langsam', 'CLK_HZ und echte Oszillatorfrequenz pruefen'),
        ('Erster Schritt zu frueh', 'Teilerreset und Resetfreigabe untersuchen'),
        ('Phase I zaehlt mehrfach', 'Entprellfenster und reales Eingangssignal messen'),
        ('Flackern / Ghosting', 'Refresh, Verdrahtung und Umschalten am Board pruefen'),
    ])),
    section("Gesamtfazit, Abgabestruktur und Quellen", [
        "Phase I und II bilden eine nachvollziehbare Entwicklung vom Segmentmuster zur zeitgesteuerten Anzeige. Das manuelle Ereignis wird durch ein Sekunden-Enable ersetzt, waehrend die gepruefte Anzeige-Basis erhalten bleibt. Beide neuen Top-Level besitzen gezielte selbstpruefende Testbenches und eigenstaendige Projektdateien.",
        "Die Abgabe enthaelt fuer beide Phasen VHDL-Quellen, Constraints, Vivado-Tcl und ein PowerShell-Simulationsskript. Dazu kommen dieser kombinierte Bericht, eine eigenstaendige Phase-I-Dokumentation und zwei Praesentationen mit Sprechernotizen. Phase III liegt als identische Kopie von prj_final unter phase_3 und wurde nicht umbenannt oder funktional angepasst.",
        "Verbleibende Arbeit vor der Laborabgabe: Autor/Gruppe ergaenzen, Synthese und Timing mit Vivado durchfuehren, Board-Prueflisten abarbeiten und die realen Messwerte eintragen. Die vorhandenen Simulationsergebnisse duerfen dabei als RTL-Nachweis verwendet werden, aber nicht als Ersatz fuer einen Hardwaretest.",
        "Anhang: src/phase2_top.vhd enthaelt den RTL-Kern, sim/tb_phase2_top.vhd die komplette automatisierte Pruefung, constraints/arty_a7_35_ssd.xdc die Pins. Die vollstaendigen Dateien liegen absichtlich direkt in der Abgabe und sind damit ohne Kopieren aus einem Bericht kompilierbar. Zur Reproduktion der Office-Dokumente dienen ../scripts/create_documents.py und requirements.txt.",
        *SOURCES,
    ]),
]


def make_doc(path, title, sections):
    doc = Document()
    setup = doc.sections[0]
    setup.page_height, setup.page_width = Cm(29.7), Cm(21)
    setup.top_margin, setup.bottom_margin = Cm(1.65), Cm(1.65)
    setup.left_margin, setup.right_margin = Cm(2), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', DocPt(10)
    normal.paragraph_format.space_after = DocPt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = DocColor.from_string(RED)
    setup.header.paragraphs[0].text = f"DIGITALE SCHALTUNGSTECHNIK  |  {title}"
    setup.header.paragraphs[0].style = doc.styles['Caption']
    footer = setup.footer.paragraphs[0]
    footer.text = 'Arty A7-35T / Pmod SSD  |  RTL-Simulation geprueft; Boardtest offen  |  Seite '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for index, item in enumerate(sections, 1):
        if index > 1:
            doc.add_page_break()
        if index == 1:
            doc.add_heading(title, 0)
            doc.add_paragraph('Laborbericht | Autor / Gruppe: noch zu ergaenzen', 'Subtitle')
        doc.add_heading(f"{index:02d}  {item['title']}", 1)
        for text in item['paragraphs']:
            paragraph = doc.add_paragraph(text)
            if text.startswith(('Digilent', 'GHDL-', 'Aufgabenstellung:', 'Lokale Referenz:')):
                for run in paragraph.runs:
                    run.font.size = DocPt(8)
        if item['image']:
            doc.add_picture(str(path.parent / item['image']), width=Cm(16.5))
        if item['table']:
            headers, rows = item['table']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Light Shading Accent 1'
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
            for row in rows:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = DocPt(3)
                        for run in paragraph.runs:
                            run.font.size = DocPt(8.5)
        if item['code']:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = DocPt(4)
            run = paragraph.add_run(item['code'])
            run.font.name, run.font.size = 'Consolas', DocPt(8)
    doc.save(path)


def make_markdown(path, title, sections):
    lines = [f'# {title}', '', 'Autor / Gruppe: noch zu ergaenzen.', '',
             'Status: RTL-Simulation bestanden; Vivado-Build und physischer Boardtest offen.', '']
    for index, item in enumerate(sections, 1):
        lines += [f"## {index:02d} {item['title']}", '', *sum(([text, ''] for text in item['paragraphs']), [])]
        if item['image']:
            lines += [f"![{item['title']}]({item['image']})", '']
        if item['table']:
            headers, rows = item['table']
            lines += ['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |']
            lines += ['| ' + ' | '.join(row) + ' |' for row in rows]
            lines += ['']
        if item['code']:
            lines += ['```text', item['code'], '```', '']
    path.write_text('\n'.join(lines), encoding='utf-8')


def textbox(slide, left, top, width, height, text, size=24, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(text.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = 'Calibri'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(12)
    return box


def slide_data(phase):
    if phase == 1:
        return [
            ('7-Segment-Anzeige', ['Phase I | Manueller Zaehler am Arty A7-35T', '0 ... 9 mit BTN0 | Reset auf 0', 'Autor / Gruppe: noch zu ergaenzen'], 'block_phase1.png',
             'Ziel vorstellen: sichtbare Ausgabe als Bruecke zwischen VHDL und realer Hardware. Die Aufgabe verlangt zunaechst 0; unser Tasterzaehler erweitert diese Basis. Keine reale Boardmessung behaupten.'),
            ('Aufgabe und Abnahmekriterien', ['Startwert rechts 0, linke Ziffer dunkel', 'Ein Tastendruck ergibt genau einen Schritt', '9 -> 0; Reset funktioniert in jedem Zustand'], None,
             'Die Anforderungen von den Aufgabenbildern abgrenzen: Ansteuerung und Testbench sind Pflicht, der Tasterzaehler ist die gewuenschte Erweiterung. Erklaeren, warum langes Halten nicht automatisch weiterzaehlen soll.'),
            ('Arty + externes Pmod SSD', ['100-MHz-Systemtakt; externe Anzeige', 'JA/JB, jeweils untere Signalreihe', '3,3 V, gemeinsame Masse, Orientierung pruefen'], None,
             'Das Arty besitzt fuer dieses Projekt keine eigene zweistellige Anzeige. Segmentkabel und Versorgung getrennt erlaeutern. Niemals unter Spannung umstecken. Pins aus der Dokumentation zeigen.'),
            ('Welche Segmente leuchten?', ['Aktiv High: 1 schaltet das Segment ein', 'Vektorreihenfolge gfedcba', '0: 0111111 | 1: 0000110'], 'segments.png',
             'Anhand der Grafik die Segmente a bis g erklaeren. Die Ziffer 1 braucht nur b und c. Die Bitreihenfolge mit den Ports seg_r[0] bis seg_r[6] verbinden und typische Verwechslungen erwaehnen.'),
            ('Architektur statt Einzellogik', ['Taster -> Synchronisation -> Entprellung', 'Zaehler -> Multiplexer -> Decoder', 'Ein Systemtakt, klar getrennte Aufgaben'], 'block_phase1.png',
             'Den Weg eines Tastendrucks bis zur rechten Ziffer verfolgen. Die Resetleitung betrifft alle Zustaende. ssd_driver und Decoder sind aus dem Weckerprojekt uebernommen und werden in Phase II wiederverwendet.'),
            ('Ein Druck, mehrere elektrische Flanken', ['Zwei Flipflops synchronisieren BTN0', '10 ms stabiler Pegel vor der Uebernahme', 'Nur akzeptierte Low->High-Flanke zaehlt'], None,
             'Metastabilitaet und mechanisches Prellen unterscheiden. Zwei Flipflops loesen das Prellproblem nicht. Das Stabilitaetsfenster wird bei Rueckkehr zum alten Zustand neu begonnen. Millionen Takte sind bei 100 MHz nur 10 ms.'),
            ('Halten ist kein zweiter Tastendruck', ['stable=0 -> 1: Zaehler +1', 'stable=1: keine Tastenwiederholung', 'Loslassen gibt den naechsten Druck frei'], None,
             'Den Zustand stable anhand Druecken-Halten-Loslassen erklaeren. Gegenueber der Wecker-Entprellung wurde Auto-Repeat bewusst nicht uebernommen. Beim Reset werden Sync-, Entprell- und Zaehlerzustaende geloescht.'),
            ('RTL: 9 ist die Umlaufgrenze', ['Begrenzter Zustand: natural range 0 to 9', 'Am gueltigen Ereignis: 9 -> 0, sonst +1', 'Keine Verzoegerungen im synthetisierbaren RTL'], None,
             'src/phase1_top.vhd als lesbaren Kern zeigen. Vergleiche verwenden den alten Signalzustand. Das Segmentmuster entsteht kombinatorisch aus count. Eine 4-Bit-Darstellung allein wuerde 10 bis 15 erlauben und waere noch kein Dezimalzaehler.'),
            ('Zwei Stellen, gemeinsame Leitungen', ['1 ms je Slot; 2 ms je vollstaendigem Bild', '500 Hz Wiederholung pro Ziffer', 'Phase I: linker Slot wird dunkel geschaltet'], None,
             'DIV=100 MHz/(2*500 Hz)=100000 rechnen. Die 500 Hz sind die Bildrate je Ziffer, nicht die Sekundenbasis. Helligkeit und moegliches Ghosting sind am Board zu pruefen und werden nicht durch die RTL-Simulation bewiesen.'),
            ('Pinbelegung und Projektanlage', ['Segmente a..d: JA7..JA10', 'Segmente e..g + Select: JB7..JB10', 'XDC + Tcl: eigenstaendiges VHDL-2008-Projekt'], None,
             'Auf die vollstaendige FPGA-Pintabelle im Bericht verweisen. E3 ist der Takt, C2 Reset, D9 BTN0. Das Skript waehlt phase1_top und tb_phase1_top. Vivado-Build ist noch offen; keine Timingwerte angeben.'),
            ('Testbench als automatischer Pruefer', ['Prellen, Halten, Loslassen, alle Ziffern', 'Rechter Slot korrekt; linker Slot dunkel', 'Reset und 9 -> 0 werden geprueft'], None,
             'Die Testbench beobachtet die externen Segmentausgaenge mit unabhaengigen Sollmustern. Bei beschleunigten Parametern werden beide Slots erreicht. Simulation an fallenden Flanken verhindert Vergleiche vor Abschluss der Delta-Zyklen.'),
            ('Ergebnis und noch offene Nachweise', ['GHDL 6.0.0: alle Phase-I-Assertions bestanden', 'Testende nach 7.290 ns', 'Vivado-Timing und realer Boardtest: offen'], None,
             'Die gemeldete Zeit ist Simulationszeit mit Testparametern und keine gemessene Reaktionszeit am Board. Optional GHDL erneut ausfuehren. Boardcheckliste und offene Timinganalyse sachlich benennen.'),
            ('Phase II baut darauf auf', ['Tasterereignis wird Sekunden-Enable', 'Zaehlerbereich waechst auf 00 ... 99', 'Decoder, Pins und Multiplexer bleiben gleich'], 'block_phase2.png',
             'Die Gesamtarchitektur fuer Phase I und II zusammenfuehren. Der neue Zaehler besitzt keinen Start-/Stop-Taster. Die Anzeige wird kontinuierlich gemultiplext, auch wenn der Sekundenzaehler zwischen zwei Ereignissen unveraendert bleibt.'),
            ('Fazit und Fragerunde', ['Eingang, Zustand und Anzeige sauber getrennt', 'Reproduzierbare Quellen und Tests enthalten', 'Quellen: Aufgabenbilder, Digilent, GHDL'], None,
             'Die drei Lernziele zusammenfassen: sichere Eingangserfassung, Zustandslogik und physikalische Anzeige. Auf ausfuehrliche Quellen-URLs im Bericht hinweisen. Nach 20 Minuten die separate Fragerunde beginnen.'),
        ]
    return [
        ('Sekundenzaehler 00 bis 99', ['Phase II | Arty A7-35T + Pmod SSD', 'Automatischer Umlauf, Reset per Taster', 'Autor / Gruppe: noch zu ergaenzen'], 'block_phase2.png',
         'Phase II kurz einordnen: Die Anzeige aus Phase I wird wiederverwendet, die Ereignisquelle wird automatisch. Der kombinierte Bericht deckt Phase I und II ab. Der Wecker ist eine separate dritte Phase.'),
        ('Was die Aufgabe fordert', ['Ein Schritt pro Sekunde von 00 bis 99', 'Nach 99 wieder 00', 'Reset setzt Zahlenstand und Zeitbasis zurueck'], None,
         'image5.png als Anforderungsquelle nennen. Fuehrende Null ist eine bewusste Darstellungsentscheidung. Der Zaehler laeuft nach Resetfreigabe ohne Starttaste. Reset darf keinen Rest der alten Sekundenphase behalten.'),
        ('Gepruefte Basis aus Phase I', ['Aktiv-High-Segmente in gfedcba-Reihenfolge', 'JA/JB: identische Verdrahtung', 'ssd_driver und Decoder unveraendert'], 'segments.png',
         'Kurz die Grundlagen wiederholen, aber nicht die gesamte erste Praesentation wiederholen. Klarstellen, dass die Basismodule lokale Kopien sind und kein Zugriff auf prj_final noetig ist.'),
        ('Ein gemeinsamer Systemtakt', ['100 MHz takten alle Register', 'Sekundenereignis ist eine Enable-Bedingung', 'Kein abgeleiteter 1-Hz-Logiktakt'], 'block_phase2.png',
         'Die Architektur erklaeren und zwischen Clock-Enable und neuer Taktdomaene unterscheiden. Der Display-Refresh laeuft unabhaengig von der Sekundenbedingung weiter. Ein langsamer Zaehler bedeutet keine langsame Anzeige.'),
        ('Von 100 MHz zu einer Sekunde', ['100.000.000 Takte pro Schritt', 'Teilerbereich 0 ... 99.999.999', 'T_step = 100.000.000 x 10 ns = 1 s'], None,
         'Die Endwertrechnung zeigen. Den klassischen Off-by-one-Fehler ansprechen: Ein Bereich bis 100 Millionen haette einen Takt zu viel. 27 Bits reichen fuer den Teiler. Die reale Genauigkeit haengt am Oszillator.'),
        ('Reset startet ein neues Intervall', ['ck_rst ist aktiv Low', 'Asynchron anfordern, synchron freigeben', 'Zahl und Teiler gemeinsam auf 0'], None,
         'Resetfreigabe ueber zwei Register erlaeutern. Ein Reset nach halber Sekunde soll nicht nach der anderen Haelfte schon weiterzaehlen. Das naechste Inkrement folgt nach einem vollen neuen Intervall. Der Fachzustand reagiert am Systemtakt.'),
        ('00 ... 99 als klarer Zustandsraum', ['count: natural range 0 to 99', 'Am Sekundenereignis: +1 oder 99 -> 0', 'Kein ungueltiger Zustand 100'], None,
         'Die Grenzen sind Teil des Entwurfs. Reset dominiert die Endwertbedingung. Die Entscheidung im Prozess liest den bisherigen Zaehlerwert, bevor die Signalzuweisung nach der Flanke wirksam wird.'),
        ('Eine Zahl, zwei Dezimalstellen', ['Einer = count mod 10', 'Zehner = count / 10', '09 -> 10 aktualisiert beide Stellen'], None,
         'Beispiele 9, 10, 59 und 99 gemeinsam durchgehen. Division durch die Konstante Zehn ist kombinatorische Logik. Beide Stellen stammen aus demselben Zahlenregister und gehen an den bekannten SSD-Treiber.'),
        ('Anzeige und Zeitbasis sind unabhaengig', ['1-ms-Slots fuer die Anzeige', '1-s-Intervall fuer den Zaehler', 'Beide Ziffern werden kontinuierlich erneuert'], None,
         'Refresh und Zaehlerfrequenz unterscheiden. Die Anzeige hat ein 2-ms-Vollbild, waehrend count eine Sekunde konstant bleibt. Deshalb muss der Anzeigetreiber auch ohne Zaehlerwechsel weiterlaufen.'),
        ('101 Sekunden Logik schnell pruefen', ['Testparameter: 1.000 Takte pro Schritt', '10-ns-Testtakt; keine reale Sekunde im Test', 'Sollwert wird an jeder Taktgrenze berechnet'], None,
         'Beschleunigung begruenden: zehn Milliarden Takte fuer den echten Umlauf waeren unnoetig teuer. Die Testbench berechnet elapsed/1000 mod 100 und prueft das Segmentmuster des aktuellen Slots. Die kleinen Parameter aendern nicht die Umlaufregel.'),
        ('Uebertraege und Reset gezielt getestet', ['Alle Zustaende 00 bis 99', '09 -> 10 und 99 -> 00', 'Reset mitten im Intervall; voller Neustart'], None,
         'Die Laufstrecke von 101357 Takten und die zweite mit 2000 Takten erklaeren. Beobachtungsflags pruefen beide Stellen. Ein reiner Zaehlerwertvergleich wuerde einen defekten Decoder nicht erkennen; hier werden die Ausgaenge geprueft.'),
        ('Simulationsergebnis und Grenzen', ['GHDL 6.0.0: alle Assertions bestanden', 'Testende nach 1.033.690 ns', 'Synthese, Timing und Boardmessung bleiben offen'], None,
         'Ausgabe aus sim/testergebnis.txt zeigen. Weder elektrische LED-Helligkeit noch Oszillatortoleranz folgen aus einer RTL-Simulation. Vivado ist hier nicht verfuegbar; es werden keine erfundenen Ressourcenwerte praesentiert.'),
        ('Boardabnahme in sinnvoller Reihenfolge', ['Zuerst Phase I: Verdrahtung und Muster', 'Dann Phase II: 00, 09 -> 10, 99 -> 00', 'Reset halten und Intervallneustart beobachten'], None,
         'Nur mit angeschlossenem Board eine echte Demonstration durchfuehren. Fuer 99->00 ausreichend Beobachtungszeit einplanen. Datum, Bitstream, Board und Messverfahren dokumentieren. Bei Problemen eine Hypothese nach der anderen pruefen.'),
        ('Vom Tasterzaehler zum Wecker', ['Phase I: manuelles Ereignis', 'Phase II: Sekunden-Enable und zwei Stellen', 'Phase III: unveraenderte Weckerkopie'], None,
         'Die Entwicklung zusammenfassen. Phase III hat 29 hashgleich kopierte Dateien; das ist ein Integritaetsnachweis und kein neuer Funktionstest. Quellen sind Aufgabenbilder, Digilent-Dokumentation und GHDL, mit URLs im Bericht. Anschliessend Fragerunde.'),
    ]


def make_ppt(path, phase):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    data = slide_data(phase)
    timings = [1, 1.5, 1.5, 1.5, 2, 1.5, 1.5, 1.5, 1.5, 1.5, 1.5, 1.5, 1, 1]
    assert math.isclose(sum(timings), 20)
    for index, (title, bullets, image, notes) in enumerate(data):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), Inches(7.5))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(RED)
        stripe.line.fill.background()
        textbox(slide, 0.65, 0.35, 11.6, 0.4, f"DIGITALE SCHALTUNGSTECHNIK / PHASE {phase}", 13, TEAL, True)
        textbox(slide, 0.65, 1.0, 11.9, 1.0, title, 34, INK, True)
        textbox(slide, 0.7, 2.2, 11.6, 2.05 if image else 3.7, '\n'.join(bullets), 23 if image else 27)
        if image:
            slide.shapes.add_picture(str(path.parent / image), Inches(0.85), Inches(4.25), width=Inches(11.6), height=Inches(2.65))
        textbox(slide, 0.7, 7.08, 10.8, 0.25, 'Arty A7-35T + Pmod SSD | Laborpraesentation', 10, GRAY)
        textbox(slide, 12.0, 7.03, 0.65, 0.35, f'{index+1:02d}', 14, RED, True)
        slide.notes_slide.notes_text_frame.text = f"Zeitbudget: {timings[index]} min; insgesamt 20 min plus Fragerunde.\n\n{notes}"
    presentation.save(path)


def main():
    for phase, sections in [(1, PHASE1), (2, PHASE1 + PHASE2)]:
        directory = ROOT / f'phase_{phase}' / 'doc'
        directory.mkdir(parents=True, exist_ok=True)
        diagram(directory / 'block_phase1.png', 1)
        diagram(directory / 'block_phase2.png', 2)
        segment_image(directory / 'segments.png')
        title = 'Phase I - 7-Segment-Anzeige' if phase == 1 else 'Phase I + II - Sekundenzaehler 00 bis 99'
        stem = 'bericht_phase_1' if phase == 1 else 'bericht_phase_1_und_2'
        make_doc(directory / f'{stem}.docx', title, sections)
        make_markdown(directory / f'{stem}.md', title, sections)
        make_ppt(directory / f'praesentation_phase_{phase}.pptx', phase)
        print(f'Phase {phase}: {len(sections)} Berichtsabschnitte, 14 Folien mit Notizen, DOCX/MD/PPTX erzeugt')


if __name__ == '__main__':
    main()